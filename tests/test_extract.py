import importlib.util
import re
from unittest import mock, skipUnless

from kb import extract
from kb.util import KBError, excerpt_in_text
from tests.helpers import VaultTestCase

HAS = {m: importlib.util.find_spec(m) is not None for m in ("trafilatura", "pymupdf", "markitdown", "pptx")}

ARTICLE = ("<html><head><title>Field Notes on Example Widgets</title>"
           "<meta name='author' content='A. Writer'></head><body><nav>Home | About</nav><article>"
           "<h1>Field Notes</h1>" + "<p>Example widgets are made of plain sample material and tested carefully "
           "by the made-up example lab before shipping to customers.</p>" * 6 +
           "<h2>Results</h2><p>The widgets passed the <a href='https://example.org/test'>drop test</a> twice.</p>"
           "</article></body></html>")


def make_pdf(path, pages=("The quick brown fox jumps over the lazy dog on page one.", ""), title="Fox Report"):
    import pymupdf
    doc = pymupdf.open()
    for t in pages:
        p = doc.new_page()
        if t:
            p.insert_text((72, 72), t)
    doc.set_metadata({"title": title, "creationDate": "D:20250301000000"})
    doc.save(path)


class TestExtract(VaultTestCase):
    @skipUnless(HAS["trafilatura"], "trafilatura not installed")
    def test_html_markdown_demoted(self):
        p = self.tmp / "page.html"
        p.write_text(ARTICLE)
        text, meta = extract.extract(p)
        self.assertTrue(meta["extracted_with"].startswith("trafilatura "))
        self.assertTrue(meta["title"].startswith("Field Notes"))
        self.assertFalse(re.search(r"(?m)^#{1,2} ", text), text)
        self.assertIn("### Results", text)
        self.assertIn("[drop test](https://example.org/test)", text)
        self.assertEqual(excerpt_in_text("made of plain sample material and tested carefully", text), "exact")

    @skipUnless(HAS["pymupdf"], "pymupdf not installed")
    def test_pdf_pages_and_ocr_note(self):
        p = self.tmp / "fox.pdf"
        make_pdf(p)
        with mock.patch("kb.ocr.available", return_value=False):
            text, meta = extract.extract(p)
        self.assertIn("[p1]\nThe quick brown fox", text)
        self.assertIn("[p2]", text)
        self.assertEqual(meta["needs_ocr"], [2])
        self.assertEqual(meta["text_quality"], "partial")
        self.assertEqual((meta["title"], meta["issued"]), ("Fox Report", "2025-03-01"))
        with mock.patch("kb.ocr.available", return_value=True), \
                mock.patch("kb.ocr.ocr_image", return_value=("scanned words here", [])):
            text, meta = extract.extract(p)
        self.assertIn("[p2]\nscanned words here", text)
        self.assertEqual((meta["text_quality"], meta["ocr_pages"]), ("mixed", [2]))

    def test_code_line_anchors(self):
        p = self.tmp / "tool.py"
        p.write_text("\n".join(f"# line {i}" for i in range(1, 121)))
        text, meta = extract.extract(p)
        self.assertEqual(re.findall(r"(?m)^\[L\d+\]$", text), ["[L1]", "[L51]", "[L101]"])
        self.assertEqual(meta["anchors"], "lines")
        self.assertEqual(excerpt_in_text("# line 51", text), "exact")
        outside = re.sub(r"(?s)```.*?```", "", text)
        self.assertNotIn("#", outside)

    def test_markdown_and_text(self):
        p = self.tmp / "notes.md"
        p.write_text("# Top\n\n## Sub\n\nbody\n\n```\n# not a heading\n```\n")
        text, _ = extract.extract(p)
        self.assertIn("### Top\n\n#### Sub", text)
        self.assertIn("# not a heading", text)

    def test_subtitles_dedupe_and_paragraphs(self):
        vtt = ("WEBVTT\n\n00:00:01.000 --> 00:00:03.000\nhello and welcome\n\n"
               "00:00:03.000 --> 00:00:05.000\nhello and welcome\nto the example talk\n\n"
               "00:00:05.000 --> 00:00:07.000\n<00:00:05.500><c>to the example talk</c>\n\n"
               "00:00:50.000 --> 00:00:52.000\nsecond part starts\n\n"
               "01:00:00.000 --> 01:00:02.000\nan hour in\n")
        p = self.tmp / "talk.en.vtt"
        p.write_text(vtt)
        text, meta = extract.extract(p)
        self.assertEqual(text, "[00:00:01] hello and welcome to the example talk\n\n[00:00:50] second part starts"
                               "\n\n[01:00:00] an hour in\n")
        srt = "1\n00:00:02,000 --> 00:00:04,000\nfirst line\n\n2\n00:00:04,000 --> 00:00:06,000\nsecond line\n"
        (self.tmp / "t.srt").write_text(srt)
        self.assertEqual(extract.extract(self.tmp / "t.srt")[0], "[00:00:02] first line second line\n")

    @skipUnless(HAS["markitdown"] and HAS["pptx"], "markitdown/python-pptx not installed")
    def test_pptx_slides(self):
        from pptx import Presentation
        prs = Presentation()
        for t in ("Alpha", "Beta"):
            s = prs.slides.add_slide(prs.slide_layouts[1])
            s.shapes.title.text = t
            s.placeholders[1].text = f"body {t}"
        p = self.tmp / "deck.pptx"
        prs.save(p)
        text, meta = extract.extract(p)
        self.assertIn("[slide 1]", text)
        self.assertIn("[slide 2]", text)
        self.assertEqual(meta["anchors"], "slides")
        self.assertFalse(re.search(r"(?m)^#{1,2} ", text), text)

    def test_image_uses_ocr(self):
        p = self.tmp / "shot.png"
        p.write_bytes(b"\x89PNG\r\n")
        with mock.patch("kb.ocr.ocr_image", return_value=("OCR TEXT", [])):
            text, meta = extract.extract(p)
        self.assertEqual((text, meta["text_quality"]), ("OCR TEXT\n", "ocr"))

    def test_missing_dependency_hint(self):
        p = self.tmp / "page.html"
        p.write_text("<html></html>")
        with mock.patch("importlib.import_module", side_effect=ImportError):
            with self.assertRaisesRegex(KBError, "pip install trafilatura"):
                extract.extract(p)

    def test_unknown_binary(self):
        p = self.tmp / "blob.bin"
        p.write_bytes(b"\xff\xfe\x00\x81" * 10)
        with self.assertRaises(KBError):
            extract.extract(p)
